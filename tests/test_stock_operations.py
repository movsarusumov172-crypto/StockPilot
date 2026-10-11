"""Write-offs, count reconciliation, returns, and SKU history."""

import csv
from collections.abc import Iterator
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from stockpilot.counting import apply_prepared_count, export_count, prepare_count
from stockpilot.database import create_database_engine
from stockpilot.models import Batch, InventoryCountLine, Product
from stockpilot.reports import (
    export_excel,
    history_report,
    inventory_report,
    loss_report,
    products_report,
)
from stockpilot.services import (
    InventoryError,
    WriteOffReason,
    receive,
    return_shipment,
    ship_with_reference,
    transfer,
    write_off,
)


@pytest.fixture
def session() -> Iterator[Session]:
    engine = create_database_engine("sqlite://")
    with Session(engine, expire_on_commit=False) as database_session:
        yield database_session
    engine.dispose()


def _receive(
    session: Session,
    sku: str,
    quantity: int,
    cost: str,
    received_on: date,
    warehouse: str = "Main",
) -> None:
    receive(
        session,
        sku=sku,
        product_name=f"Product {sku}",
        warehouse_name=warehouse,
        received_on=received_on,
        quantity=quantity,
        unit_cost=Decimal(cost),
    )


def test_write_off_uses_fifo_and_is_reported_as_a_loss(session: Session) -> None:
    with session.begin():
        _receive(session, "DEF-1", 2, "2.00", date(2026, 1, 1))
        _receive(session, "DEF-1", 3, "5.00", date(2026, 1, 2))
    with session.begin():
        amount = write_off(
            session,
            sku="DEF-1",
            warehouse_name="Main",
            occurred_on=date(2026, 1, 3),
            quantity=4,
            reason=WriteOffReason.DEFECT,
        )

    losses = loss_report(session)
    assert amount == Decimal("14.00")
    assert losses["quantity"].sum() == 4
    assert losses["amount"].sum() == Decimal("14.00")
    assert set(losses["reason"]) == {WriteOffReason.DEFECT.value}


def test_count_applies_both_surplus_and_fifo_shortage(session: Session, tmp_path: Path) -> None:
    with session.begin():
        _receive(session, "SHORT-1", 2, "2.00", date(2026, 1, 1))
        _receive(session, "SHORT-1", 3, "5.00", date(2026, 1, 2))
        _receive(session, "SURPLUS-1", 4, "3.00", date(2026, 1, 2))

    count_file = tmp_path / "count.csv"
    export_count(session, "Main", count_file)
    with count_file.open("r", encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.DictReader(csv_file))
        fieldnames = rows[0].keys()
    for row in rows:
        row["counted_quantity"] = "2" if row["sku"] == "SHORT-1" else "6"
    with count_file.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    session.rollback()
    with session.begin():
        warehouse, previews = prepare_count(session, count_file)
        amounts = {preview.adjustment.sku: preview.amount for preview in previews}
        assert amounts == {"SHORT-1": Decimal("9.00"), "SURPLUS-1": Decimal("6.00")}
        apply_prepared_count(session, warehouse, previews, date(2026, 1, 5))

    products = session.scalars(select(Product)).all()
    balances = {
        product.sku: sum(
            batch.quantity_remaining
            for batch in session.scalars(select(Batch).where(Batch.product_id == product.id)).all()
        )
        for product in products
    }
    assert balances == {"SHORT-1": 2, "SURPLUS-1": 6}

    count_lines = session.scalars(select(InventoryCountLine)).all()
    assert {line.difference for line in count_lines} == {-3, 2}
    shortages = loss_report(session)
    assert shortages["amount"].sum() == Decimal("9.00")


def test_count_rejects_a_stale_export(session: Session, tmp_path: Path) -> None:
    with session.begin():
        _receive(session, "STALE-1", 3, "1.50", date(2026, 1, 1))
    count_file = tmp_path / "count.csv"
    export_count(session, "Main", count_file)
    with count_file.open("r", encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.DictReader(csv_file))
        fieldnames = rows[0].keys()
    rows[0]["counted_quantity"] = "3"
    with count_file.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    session.rollback()

    with session.begin():
        receive(
            session,
            sku="STALE-1",
            product_name="Product STALE-1",
            warehouse_name="Main",
            received_on=date(2026, 1, 2),
            quantity=1,
            unit_cost=Decimal("1.50"),
        )
    session.rollback()

    with pytest.raises(InventoryError, match="изменился после выгрузки"):
        with session.begin():
            prepare_count(session, count_file)


def test_return_restores_original_shipment_cost_and_limits_quantity(session: Session) -> None:
    with session.begin():
        _receive(session, "RET-1", 2, "2.00", date(2026, 1, 1))
        _receive(session, "RET-1", 3, "5.00", date(2026, 1, 2))
    with session.begin():
        shipment = ship_with_reference(
            session,
            sku="RET-1",
            warehouse_name="Main",
            shipped_on=date(2026, 1, 3),
            quantity=4,
        )
    assert shipment.cogs == Decimal("14.00")

    with session.begin():
        returned_cost = return_shipment(
            session,
            shipment_id=shipment.id,
            returned_on=date(2026, 1, 4),
            quantity=3,
        )
    assert returned_cost == Decimal("9.00")

    with pytest.raises(InventoryError, match="only 1 units available to return"):
        with session.begin():
            return_shipment(
                session,
                shipment_id=shipment.id,
                returned_on=date(2026, 1, 5),
                quantity=2,
            )

    return_batches = session.scalars(
        select(Batch).where(Batch.received_on == date(2026, 1, 4))
    ).all()
    assert sorted(batch.unit_cost for batch in return_batches) == [
        Decimal("2.00"),
        Decimal("5.00"),
    ]


def test_history_tracks_balance_after_each_transfer_movement(session: Session) -> None:
    with session.begin():
        _receive(session, "HIST-1", 3, "2.00", date(2026, 1, 1))
    with session.begin():
        transfer(
            session,
            sku="HIST-1",
            source_name="Main",
            destination_name="North",
            transferred_on=date(2026, 1, 2),
            quantity=1,
        )

    history = history_report(session, "HIST-1")
    assert list(history["balance_after"]) == [3, 2, 1]
    assert list(history["kind"]) == ["receive", "transfer_out", "transfer_in"]


def test_product_search_lists_balances_by_warehouse(session: Session) -> None:
    with session.begin():
        _receive(session, "LOOK-1", 3, "2.00", date(2026, 1, 1), "Main")
        _receive(session, "LOOK-1", 4, "2.50", date(2026, 1, 2), "North")
        _receive(session, "OTHER-1", 8, "1.00", date(2026, 1, 2), "Main")

    products = products_report(session, "LOOK")
    main_inventory = inventory_report(session, "Main")

    assert set(products["sku"]) == {"LOOK-1"}
    assert dict(zip(products["warehouse"], products["quantity"], strict=True)) == {
        "Main": 3,
        "North": 4,
    }
    assert set(main_inventory["warehouse"]) == {"Main"}


def test_loss_report_can_filter_a_warehouse(session: Session) -> None:
    with session.begin():
        _receive(session, "LOSS-1", 2, "3.00", date(2026, 1, 1), "Main")
        _receive(session, "LOSS-1", 2, "5.00", date(2026, 1, 1), "North")
        write_off(
            session,
            sku="LOSS-1",
            warehouse_name="Main",
            occurred_on=date(2026, 1, 2),
            quantity=1,
            reason=WriteOffReason.DEFECT,
        )
        write_off(
            session,
            sku="LOSS-1",
            warehouse_name="North",
            occurred_on=date(2026, 1, 2),
            quantity=1,
            reason=WriteOffReason.SPOILAGE,
        )

    losses = loss_report(session, "North")

    assert set(losses["warehouse"]) == {"North"}
    assert losses["amount"].sum() == Decimal("5.00")


def test_excel_export_contains_a_separate_losses_sheet(session: Session, tmp_path: Path) -> None:
    with session.begin():
        _receive(session, "XLS-1", 3, "2.25", date(2026, 1, 1))
        write_off(
            session,
            sku="XLS-1",
            warehouse_name="Main",
            occurred_on=date(2026, 1, 2),
            quantity=1,
            reason=WriteOffReason.SPOILAGE,
        )
    output = tmp_path / "reports" / "stock.xlsx"

    export_excel(session, output, date(2026, 1, 1), date(2026, 1, 31))

    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=True)
    assert "Потери" in workbook.sheetnames
    assert workbook["Потери"].max_row == 2
    workbook.close()
